#!/usr/bin/env python3
"""Shared implementation for the document helper commands."""
import argparse
import io
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path


def run(*args):
    try:
        return subprocess.run([str(a) for a in args], check=True)
    except FileNotFoundError as exc:
        raise RuntimeError(f"Required tool is missing: {args[0]}") from exc


def png(data, output):
    from PIL import Image, ImageOps
    with Image.open(io.BytesIO(data)) as image:
        image = ImageOps.exif_transpose(image)
        image = image.convert("RGBA" if "A" in image.getbands() or "transparency" in image.info else "RGB")
        image.thumbnail((1600, 1600))
        image.save(output, "PNG")


def extract(source, output):
    if source.suffix.lower() == ".pdf":
        import fitz
        with fitz.open(source) as doc:
            for page in doc:
                for index, item in enumerate(page.get_images(full=True)):
                    # Render the image with its soft mask, when present.
                    pix = fitz.Pixmap(doc, item[0])
                    if item[1]:
                        pix = fitz.Pixmap(pix, fitz.Pixmap(doc, item[1]))
                    if pix.colorspace and pix.colorspace.n > 3:
                        pix = fitz.Pixmap(fitz.csRGB, pix)
                    png(pix.tobytes("png"), output / f"page-{page.number + 1}-{index + 1}.png")
    elif source.suffix.lower() in {".docx", ".pptx", ".xlsx"}:
        with zipfile.ZipFile(source) as archive:
            for index, name in enumerate(archive.namelist()):
                if "/media/" not in name or name.endswith("/"):
                    continue
                target = output / f"image-{index}.png"
                data = archive.read(name)
                if Path(name).suffix.lower() in {".emf", ".wmf"}:
                    with tempfile.TemporaryDirectory(dir=output) as temp:
                        vector = Path(temp) / Path(name).name
                        vector.write_bytes(data)
                        rendered = Path(temp) / "rendered.png"
                        run("inkscape", vector, "--export-type=png", f"--export-filename={rendered}")
                        png(rendered.read_bytes(), target)
                else:
                    png(data, target)
    else:
        png(source.read_bytes(), output / "image.png")


def text(source, output):
    kind = source.suffix.lower()
    target = output / "document.md"
    if kind == ".pdf":
        import fitz
        with fitz.open(source) as doc, target.open("w", encoding="utf-8") as result:
            for page in doc:
                with tempfile.TemporaryDirectory(dir=output) as temp:
                    page_text = Path(temp) / "page.txt"
                    run("pdftotext", "-layout", "-f", page.number + 1, "-l", page.number + 1, source, page_text)
                    content = page_text.read_text(encoding="utf-8")
                    if not content.strip():
                        raster = Path(temp) / "page.png"
                        page.get_pixmap(dpi=200).save(raster)
                        run("tesseract", raster, Path(temp) / "ocr", "-l", "eng+pol")
                        content = (Path(temp) / "ocr.txt").read_text(encoding="utf-8")
                    result.write(f"\n\n## Page {page.number + 1}\n\n{content}")
    elif kind in {".docx", ".html", ".htm"}:
        run("pandoc", source, "-t", "gfm", "--wrap=none", f"--extract-media={output / 'media'}", "-o", target)
    elif kind == ".xlsx":
        import openpyxl
        book = openpyxl.load_workbook(source, read_only=True, data_only=True)
        try:
            with target.open("w", encoding="utf-8") as result:
                for sheet in book:
                    result.write(f"\n## {sheet.title}\n\n")
                    for row in sheet.values:
                        result.write("\t".join("" if value is None else str(value) for value in row) + "\n")
        finally:
            book.close()
    elif kind == ".pptx":
        from pptx import Presentation
        with target.open("w", encoding="utf-8") as result:
            for index, slide in enumerate(Presentation(source).slides, 1):
                result.write(f"\n## Slide {index}\n\n")
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        result.write(shape.text + "\n\n")
                    if shape.has_table:
                        for row in shape.table.rows:
                            result.write("\t".join(cell.text for cell in row.cells) + "\n")
    else:
        raise RuntimeError(f"Unsupported document type: {kind}")
    print(target)


def main():
    parser = argparse.ArgumentParser(description=__doc__, epilog=(
        "doc-to-text <file> [outdir]; doc-extract-images <file> [outdir]; "
        "pdf-to-png <file> [first] [last] [dpi] [--outdir directory]. "
        "Default output: /tmp/doc-out/<basename>; default dpi: 110."))
    parser.add_argument("command", choices=["doc-to-text", "doc-extract-images", "pdf-to-png"])
    parser.add_argument("file", type=Path)
    parser.add_argument("parameters", nargs="*")
    parser.add_argument("--outdir", type=Path, help="Output directory (all commands)")
    args = parser.parse_args()
    source = args.file.resolve(strict=True)
    if args.command != "pdf-to-png" and len(args.parameters) > 1:
        parser.error("Expected at most one output directory")
    output = args.outdir or (Path(args.parameters[0]) if args.parameters and args.command != "pdf-to-png" else Path("/tmp/doc-out") / source.stem)
    output = output.resolve()
    # Never overwrite a source, even if it has a helper's output filename.
    if source == output or output in source.parents:
        parser.error("The output directory must not contain the source file")
    output.mkdir(parents=True, exist_ok=True)
    if args.command == "doc-to-text":
        text(source, output)
    elif args.command == "doc-extract-images":
        extract(source, output)
        print(output)
    else:
        import fitz
        if len(args.parameters) > 3:
            parser.error("Expected [first] [last] [dpi]")
        with fitz.open(source) as doc:
            first, last, dpi = (list(map(int, args.parameters)) + [None] * 3)[:3]
            first = 1 if first is None else first
            last = len(doc) if last is None else last
            dpi = 110 if dpi is None else dpi
            if not 1 <= first <= last <= len(doc) or not 1 <= dpi <= 600:
                parser.error("Invalid page range or dpi (1..600)")
            for number in range(first, last + 1):
                path = output / f"page-{number}.png"
                doc[number - 1].get_pixmap(dpi=dpi).save(path)
                print(path)


if __name__ == "__main__":
    try:
        main()
    except (RuntimeError, ImportError, OSError, ValueError, subprocess.CalledProcessError) as error:
        sys.exit(f"Document helper failed: {error}")
